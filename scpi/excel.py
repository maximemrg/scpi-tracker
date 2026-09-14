"""Génération de SCPI_tracker.xlsx (openpyxl) avec mise en forme.

Onglets : Dashboard, Comparatif, Portefeuille, Historique, Flux, Sources,
A_VERIFIER. Chaque cellule de donnée porte un commentaire (source + date +
confiance). Les cellules de confiance FAIBLE sont sur fond orange ; les cases
en échec d'extraction (A_VERIFIER) sont grisées avec le lien direct.

Aucune valeur n'est inventée : une case sans donnée fiable affiche « À VÉRIFIER »
et renvoie à sa source.
"""

from __future__ import annotations

from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.cell.cell import Cell as XLCell
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .models import Confidence, MetricKey, Status
from .registry import Registry
from .storage import Store

# Palette
_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_FAIBLE_FILL = PatternFill("solid", fgColor="FFD9A0")   # orange = confiance FAIBLE
_MOYENNE_FILL = PatternFill("solid", fgColor="FFF3CC")  # jaune pâle = MOYENNE
_MANUELLE_FILL = PatternFill("solid", fgColor="DDEBF7")  # bleu pâle = saisie manuelle
_AVERIF_FILL = PatternFill("solid", fgColor="F2F2F2")   # gris = à vérifier
_TITLE_FONT = Font(size=14, bold=True, color="1F3864")


@dataclass
class CellData:
    value: float | str | None
    unit: str | None
    confidence: str | None
    status: str
    source_url: str | None
    published_at: str | None


# Colonnes du Comparatif : (metric_key, libellé, unité, sens_favorable)
# sens_favorable : "haut" = vert quand élevé, "bas" = vert quand faible, None = neutre
COMPARATIF_COLS: list[tuple[str, str, str, str | None]] = [
    (MetricKey.PRIX_SOUSCRIPTION, "Prix souscription", "€", None),
    (MetricKey.PRIX_RETRAIT, "Prix retrait", "€", None),
    (MetricKey.VALEUR_RECONSTITUTION, "Val. reconstitution", "€", None),
    (MetricKey.VALEUR_REALISATION, "Val. réalisation", "€", None),
    ("ecart_prix_reconstitution", "Écart prix/reconst.", "%", "bas"),
    (MetricKey.TAUX_DISTRIBUTION, "TD", "%", "haut"),
    (MetricKey.TRI_5ANS, "TRI 5 ans", "%", "haut"),
    (MetricKey.TRI_10ANS, "TRI 10 ans", "%", "haut"),
    (MetricKey.TOF, "TOF", "%", "haut"),
    (MetricKey.CAPITALISATION, "Capitalisation", "€", None),
    (MetricKey.NOMBRE_ASSOCIES, "Associés", "", None),
    (MetricKey.NOMBRE_IMMEUBLES, "Immeubles", "", None),
    (MetricKey.DELAI_JOUISSANCE, "Délai jouiss.", "mois", "bas"),
    (MetricKey.FRAIS_GESTION, "Frais gestion", "%", "bas"),
]


def _fetch_cell(store: Store, scpi_id: str, key: str) -> CellData | None:
    """Valeur affichable : override manuel > dernière OK > dernière A_VERIFIER."""
    ov = store.conn.execute(
        """SELECT value_num, value_text, unit, source_url, published_at
           FROM manual_overrides WHERE scpi_id=? AND metric_key=?
           ORDER BY created_at DESC, id DESC LIMIT 1""",
        (scpi_id, key),
    ).fetchone()
    if ov is not None:
        return CellData(ov["value_num"] if ov["value_num"] is not None else ov["value_text"],
                    ov["unit"], Confidence.MANUELLE, Status.OK, ov["source_url"],
                    ov["published_at"])
    ok = store.conn.execute(
        """SELECT value_num, value_text, unit, confidence, source_url, published_at
           FROM metrics WHERE scpi_id=? AND metric_key=? AND status='OK'
           ORDER BY collected_at DESC, id DESC LIMIT 1""",
        (scpi_id, key),
    ).fetchone()
    if ok is not None:
        return CellData(ok["value_num"] if ok["value_num"] is not None else ok["value_text"],
                    ok["unit"], ok["confidence"], Status.OK, ok["source_url"],
                    ok["published_at"])
    av = store.conn.execute(
        """SELECT source_url FROM metrics WHERE scpi_id=? AND metric_key=? AND status='A_VERIFIER'
           ORDER BY collected_at DESC, id DESC LIMIT 1""",
        (scpi_id, key),
    ).fetchone()
    if av is not None:
        return CellData(None, None, Confidence.FAIBLE, Status.A_VERIFIER, av["source_url"], None)
    return None


def _header_row(ws: Worksheet, headers: list[str], row: int = 1) -> None:
    for col, text in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col, value=text)
        c.fill = _HEADER_FILL
        c.font = _HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _apply_value_cell(cell: XLCell, data: CellData) -> None:
    """Écrit la valeur, la couleur selon la confiance, et le commentaire source."""
    if data.status == Status.A_VERIFIER or data.value is None:
        cell.value = "À VÉRIFIER"
        cell.fill = _AVERIF_FILL
        cell.font = Font(italic=True, color="9C5700")
        if data.source_url:
            cell.comment = Comment(f"Extraction incertaine.\nSource : {data.source_url}", "SCPI")
        return
    cell.value = data.value
    conf = data.confidence
    if conf == Confidence.FAIBLE:
        cell.fill = _FAIBLE_FILL
    elif conf == Confidence.MOYENNE:
        cell.fill = _MOYENNE_FILL
    elif conf == Confidence.MANUELLE:
        cell.fill = _MANUELLE_FILL
    parts = []
    if data.published_at:
        parts.append(f"Publié : {data.published_at}")
    parts.append(f"Confiance : {conf}")
    if data.source_url:
        parts.append(f"Source : {data.source_url}")
    cell.comment = Comment("\n".join(parts), "SCPI")


def _sheet_comparatif(wb: Workbook, store: Store) -> None:
    ws = wb.active
    ws.title = "Comparatif"
    headers = ["SCPI", "Société de gestion"] + [c[1] for c in COMPARATIF_COLS]
    _header_row(ws, headers)
    scpis = store.conn.execute("SELECT * FROM scpi ORDER BY sdg_nom, nom").fetchall()
    r = 2
    for s in scpis:
        ws.cell(row=r, column=1, value=s["nom"]).font = Font(bold=True)
        ws.cell(row=r, column=2, value=s["sdg_nom"])
        cells = {k: _fetch_cell(store, s["scpi_id"], k) for k, *_ in COMPARATIF_COLS
                 if k != "ecart_prix_reconstitution"}
        # Écart prix/reconstitution dérivé (si les 2 valeurs sont disponibles).
        ps = cells.get(MetricKey.PRIX_SOUSCRIPTION)
        vr = cells.get(MetricKey.VALEUR_RECONSTITUTION)
        ecart: CellData | None = None
        if (ps and vr and isinstance(ps.value, (int, float))
                and isinstance(vr.value, (int, float)) and vr.value):
            pct = (ps.value - vr.value) / vr.value * 100
            ecart = CellData(round(pct, 2), "%", Confidence.MOYENNE, Status.OK, None,
                         ps.published_at)
            ecart.source_url = "Calculé : (prix soucription − reconstitution) / reconstitution"

        for col_idx, (key, _label, _unit, _fav) in enumerate(COMPARATIF_COLS, start=3):
            data = ecart if key == "ecart_prix_reconstitution" else cells.get(key)
            cell = ws.cell(row=r, column=col_idx)
            if data is None:
                cell.value = "—"
                cell.font = Font(color="BFBFBF")
            else:
                _apply_value_cell(cell, data)
        r += 1

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{r - 1}"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 22
    for col_idx in range(3, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 14
    # Mise en forme conditionnelle (échelles vert/rouge) sur colonnes clés.
    last = r - 1
    for col_idx, (_key, _label, _unit, fav) in enumerate(COMPARATIF_COLS, start=3):
        if fav is None or last < 2:
            continue
        letter = get_column_letter(col_idx)
        rng = f"{letter}2:{letter}{last}"
        if fav == "haut":
            rule = ColorScaleRule(start_type="min", start_color="F8696B",
                                  mid_type="percentile", mid_value=50, mid_color="FFEB84",
                                  end_type="max", end_color="63BE7B")
        else:  # bas = vert
            rule = ColorScaleRule(start_type="min", start_color="63BE7B",
                                  mid_type="percentile", mid_value=50, mid_color="FFEB84",
                                  end_type="max", end_color="F8696B")
        ws.conditional_formatting.add(rng, rule)


def _sheet_sources(wb: Workbook, store: Store) -> None:
    ws = wb.create_sheet("Sources")
    _header_row(ws, ["SCPI", "Métrique", "Période", "Valeur", "Unité",
                     "Date publication", "Confiance", "Source (URL)"])
    rows = store.conn.execute(
        """SELECT scpi_id, metric_key, period, value_num, value_text, unit,
                  published_at, confidence, source_url
           FROM metrics WHERE status='OK'
           ORDER BY scpi_id, metric_key, collected_at DESC"""
    ).fetchall()
    r = 2
    seen: set[tuple[str, str, str | None]] = set()
    for row in rows:
        keyt = (row["scpi_id"], row["metric_key"], row["period"])
        if keyt in seen:
            continue
        seen.add(keyt)
        val = row["value_num"] if row["value_num"] is not None else row["value_text"]
        ws.cell(row=r, column=1, value=row["scpi_id"])
        ws.cell(row=r, column=2, value=row["metric_key"])
        ws.cell(row=r, column=3, value=row["period"])
        ws.cell(row=r, column=4, value=val)
        ws.cell(row=r, column=5, value=row["unit"])
        ws.cell(row=r, column=6, value=row["published_at"])
        cc = ws.cell(row=r, column=7, value=row["confidence"])
        if row["confidence"] == Confidence.FAIBLE:
            cc.fill = _FAIBLE_FILL
        ws.cell(row=r, column=8, value=row["source_url"])
        r += 1
    _autofit(ws, [24, 22, 10, 16, 8, 16, 10, 70])
    ws.auto_filter.ref = f"A1:H{max(r - 1, 1)}"
    ws.freeze_panes = "A2"


def _sheet_a_verifier(wb: Workbook, store: Store) -> None:
    ws = wb.create_sheet("A_VERIFIER")
    _header_row(ws, ["SCPI", "Métrique", "Période", "Raison", "Lien direct"])
    rows = store.conn.execute(
        """SELECT scpi_id, metric_key, period, note, source_url, MAX(collected_at)
           FROM metrics WHERE status='A_VERIFIER'
           GROUP BY scpi_id, metric_key, period
           ORDER BY scpi_id, metric_key"""
    ).fetchall()
    r = 2
    for row in rows:
        ws.cell(row=r, column=1, value=row["scpi_id"])
        ws.cell(row=r, column=2, value=row["metric_key"])
        ws.cell(row=r, column=3, value=row["period"])
        ws.cell(row=r, column=4, value=row["note"])
        ws.cell(row=r, column=5, value=row["source_url"])
        r += 1
    _autofit(ws, [24, 22, 10, 60, 70])
    ws.auto_filter.ref = f"A1:E{max(r - 1, 1)}"
    ws.freeze_panes = "A2"
    if r == 2:
        ws.cell(row=2, column=1,
                value="(vide — vérifier qu'aucune extraction ne triche !)")


def _sheet_historique(wb: Workbook, store: Store) -> None:
    ws = wb.create_sheet("Historique")
    _header_row(ws, ["SCPI", "Métrique", "Période", "Valeur", "Date publication",
                     "Date collecte", "Confiance"])
    rows = store.conn.execute(
        """SELECT scpi_id, metric_key, period, value_num, published_at, collected_at, confidence
           FROM metrics
           WHERE status='OK' AND metric_key IN (?, ?, ?)
           ORDER BY scpi_id, metric_key, collected_at""",
        (MetricKey.PRIX_SOUSCRIPTION, MetricKey.TAUX_DISTRIBUTION, MetricKey.ACOMPTE),
    ).fetchall()
    r = 2
    for row in rows:
        for i, v in enumerate([row["scpi_id"], row["metric_key"], row["period"],
                               row["value_num"], row["published_at"],
                               row["collected_at"][:10], row["confidence"]], start=1):
            ws.cell(row=r, column=i, value=v)
        r += 1
    _autofit(ws, [24, 20, 10, 14, 16, 12, 10])
    ws.auto_filter.ref = f"A1:G{max(r - 1, 1)}"
    ws.freeze_panes = "A2"


def _sheet_flux(wb: Workbook, store: Store) -> None:
    ws = wb.create_sheet("Flux")
    _header_row(ws, ["SCPI", "Période", "Acompte / part", "Date publication",
                     "Confiance", "Source"])
    rows = store.conn.execute(
        """SELECT scpi_id, period, value_num, published_at, confidence, source_url
           FROM metrics WHERE status='OK' AND metric_key=?
           ORDER BY scpi_id, period""",
        (MetricKey.ACOMPTE,),
    ).fetchall()
    r = 2
    for row in rows:
        for i, v in enumerate([row["scpi_id"], row["period"], row["value_num"],
                               row["published_at"], row["confidence"], row["source_url"]],
                              start=1):
            ws.cell(row=r, column=i, value=v)
        r += 1
    if r == 2:
        ws.cell(row=2, column=1, value="(acomptes à compléter — voir A_VERIFIER / saisie)")
    _autofit(ws, [24, 12, 14, 16, 10, 60])
    ws.freeze_panes = "A2"


def _sheet_dashboard(wb: Workbook, store: Store) -> None:
    ws = wb.create_sheet("Dashboard", 0)
    ws["A1"] = "SCPI Tracker — Tableau de bord"
    ws["A1"].font = _TITLE_FONT
    ws["A3"] = ("Note : la vue portefeuille (valeur, TRI global) sera alimentée "
               "par portefeuille.yaml.")
    ws["A3"].font = Font(italic=True, color="808080")

    n_scpi = store.conn.execute("SELECT COUNT(*) c FROM scpi").fetchone()["c"]
    n_ok = store.conn.execute("SELECT COUNT(*) c FROM metrics WHERE status='OK'").fetchone()["c"]
    n_av = store.conn.execute(
        "SELECT COUNT(*) c FROM metrics WHERE status='A_VERIFIER'").fetchone()["c"]
    stats = [("SCPI suivies", n_scpi), ("Métriques OK", n_ok), ("À vérifier", n_av)]
    for i, (label, val) in enumerate(stats, start=5):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        ws.cell(row=i, column=2, value=val)

    # TD par SCPI (pour un graphique) : dernière valeur OK connue.
    ws["A10"] = "Taux de distribution par SCPI"
    ws["A10"].font = Font(bold=True)
    _header_row(ws, ["SCPI", "TD %"], row=11)
    td_rows = store.conn.execute(
        """SELECT s.nom, m.value_num FROM scpi s
           JOIN metrics m ON m.scpi_id=s.scpi_id
           WHERE m.metric_key=? AND m.status='OK' AND m.value_num IS NOT NULL
           GROUP BY s.scpi_id HAVING MAX(m.collected_at)
           ORDER BY m.value_num DESC""",
        (MetricKey.TAUX_DISTRIBUTION,),
    ).fetchall()
    start = 12
    for i, row in enumerate(td_rows):
        ws.cell(row=start + i, column=1, value=row["nom"])
        ws.cell(row=start + i, column=2, value=row["value_num"])
    end = start + len(td_rows) - 1
    if td_rows:
        chart = BarChart()
        chart.title = "Taux de distribution (dernier connu)"
        chart.type = "bar"
        chart.height = 1.2 * len(td_rows) + 3
        data = Reference(ws, min_col=2, min_row=11, max_row=end)
        cats = Reference(ws, min_col=1, min_row=12, max_row=end)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws.add_chart(chart, "D11")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 10


def _sheet_portefeuille(wb: Workbook) -> None:
    ws = wb.create_sheet("Portefeuille")
    ws["A1"] = "Portefeuille — alimenté par portefeuille.yaml (à créer à l'étape dédiée)"
    ws["A1"].font = _TITLE_FONT
    _header_row(ws, ["SCPI", "Date achat", "Nb parts", "Prix unitaire", "Frais",
                     "Détention", "Démembrement (années)", "Support",
                     "Valeur courante", "PRU", "+/- value latente", "TRI réel"], row=3)
    _autofit(ws, [24, 12, 10, 13, 10, 16, 20, 14, 15, 12, 16, 10])
    ws.freeze_panes = "A4"


def _autofit(ws: Worksheet, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_workbook(store: Store, path: str, reg: Registry) -> None:
    wb = Workbook()
    _sheet_comparatif(wb, store)   # devient l'onglet actif renommé
    _sheet_sources(wb, store)
    _sheet_a_verifier(wb, store)
    _sheet_historique(wb, store)
    _sheet_flux(wb, store)
    _sheet_portefeuille(wb)
    _sheet_dashboard(wb, store)         # inséré en position 0
    wb.save(path)


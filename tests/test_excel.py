"""Test du générateur Excel : structure, commentaires source, override prioritaire."""

from datetime import date

import openpyxl

from scpi.excel import build_workbook
from scpi.models import Confidence, Metric, MetricKey, now_utc
from scpi.registry import load_registry
from scpi.storage import Store


def _seed(store: Store) -> None:
    store.upsert_scpi(
        "corum_eurion", "Corum Eurion", "corum_lepargne", "Corum L'Épargne", "corum.fr"
    )
    store.insert_metrics([
        Metric.ok_num(scpi_id="corum_eurion", metric_key=MetricKey.TAUX_DISTRIBUTION,
                      value=5.73, unit="%", period="2025", source_url="https://corum.fr/x.pdf",
                      published_at=date(2026, 6, 30), collected_at=now_utc(),
                      confidence=Confidence.HAUTE),
        Metric.a_verifier(scpi_id="corum_eurion", metric_key=MetricKey.TOF,
                          source_url="https://corum.fr/x.pdf", collected_at=now_utc(),
                          note="illisible"),
    ])


def test_build_workbook(tmp_path):
    store = Store(tmp_path / "t.sqlite")
    _seed(store)
    out = tmp_path / "SCPI_tracker.xlsx"
    build_workbook(store, str(out), load_registry())
    store.close()

    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == [
        "Dashboard", "Comparatif", "Sources", "A_VERIFIER", "Historique", "Flux", "Portefeuille",
    ]
    comp = wb["Comparatif"]
    # Ligne Corum Eurion présente, TD renseigné, TOF en "À VÉRIFIER".
    row = {comp.cell(1, c).value: comp.cell(2, c).value for c in range(1, comp.max_column + 1)}
    assert row["SCPI"] == "Corum Eurion"
    assert row["TD"] == 5.73
    assert row["TOF"] == "À VÉRIFIER"
    # Le TD porte un commentaire mentionnant sa source.
    td_col = [c for c in range(1, comp.max_column + 1) if comp.cell(1, c).value == "TD"][0]
    assert "corum.fr" in comp.cell(2, td_col).comment.text


def test_override_wins_in_excel(tmp_path):
    store = Store(tmp_path / "t.sqlite")
    _seed(store)
    store.add_override(scpi_id="corum_eurion", metric_key=MetricKey.TOF, period=None,
                       value_num=97.5, unit="%", note="saisie")
    out = tmp_path / "w.xlsx"
    build_workbook(store, str(out), load_registry())
    store.close()
    comp = openpyxl.load_workbook(out)["Comparatif"]
    tof_col = [c for c in range(1, comp.max_column + 1) if comp.cell(1, c).value == "TOF"][0]
    assert comp.cell(2, tof_col).value == 97.5  # override MANUELLE gagne sur A_VERIFIER
